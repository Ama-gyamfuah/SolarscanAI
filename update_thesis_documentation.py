import docx
import shutil
import os

doc_path = "documents/Solar Scan Final Year Project Documentation.docx"
doc = docx.Document(doc_path)
print("Loaded document successfully. Total paragraphs:", len(doc.paragraphs))

# ─────────────────────────────────────────────────────────────
# 1. UPDATE SECTION 1.6 RESEARCH OBJECTIVES (EXACTLY 3 SPECIFIC OBJECTIVES)
# ─────────────────────────────────────────────────────────────
obj_start = -1
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == "1.6 Research Objectives":
        obj_start = i
        break

print(f"Found 1.6 Research Objectives at P{obj_start}")

# P+1 is 'General Objective'
# P+2 is General Objective paragraph text
# P+3 is 'Specific Objectives'
# P+4 is intro: 'The following specific objectives are set out to achieve the general objective stated above:'
# P+5, P+6, P+7, P+8, P+9 are the 5 previous specific objectives.

# Update General Objective text slightly to reflect SOLAR SCAN
gen_obj_p = doc.paragraphs[obj_start + 2]
gen_obj_p.text = (
    "The general objective of this study is to design, develop, and evaluate an automated, "
    "edge-deployed mobile system, named SOLAR SCAN, for the offline detection, classification, "
    "and physical characterization of photovoltaic panel defects. This is achieved through an "
    "optimized YOLOv8 deep learning model combined with TensorFlow Lite (TFLite) mobile integration, "
    "giving clean energy technicians in Ghana and other developing regions a reliable, low-cost "
    "diagnostic tool that operates independently of continuous internet connectivity."
)

# New 3 Consolidated Specific Objectives
specific_objs = [
    (
        "1. To curate an annotated multi-modal dataset (visual RGB, thermal infrared, and electroluminescence) "
        "and train a custom YOLOv8 object detection model, applying post-training INT8 quantization to compress "
        "model weights from 6.3 MB to 3.2 MB for rapid, offline edge inference on mobile CPUs."
    ),
    (
        "2. To design and implement the SOLAR SCAN software architecture integrating a client-side Double-Layer "
        "Validation Gatekeeper to prevent out-of-domain false inferences, coupled with a dynamic wafer thermodynamics "
        "engine that calculates cell operating temperatures and real wattage losses via NOCT physical modeling."
    ),
    (
        "3. To empirically evaluate the diagnostic accuracy, inference latency, and field usability of the complete "
        "system using benchmark detection metrics (precision, recall, mAP) alongside standardized System Usability "
        "Scale (SUS) assessments with clean energy technicians in Ghana."
    )
]

# Set P+5, P+6, P+7 to the 3 specific objectives
doc.paragraphs[obj_start + 5].text = specific_objs[0]
doc.paragraphs[obj_start + 6].text = specific_objs[1]
doc.paragraphs[obj_start + 7].text = specific_objs[2]

# Remove P+8 and P+9 (the old 4th and 5th objectives)
p8 = doc.paragraphs[obj_start + 8]
p9 = doc.paragraphs[obj_start + 9]
p8._p.getparent().remove(p8._p)
p9._p.getparent().remove(p9._p)

print("Updated Section 1.6: Specific Objectives consolidated to exactly 3.")

# ─────────────────────────────────────────────────────────────
# 2. UPDATE TABLE OF FIGURES (PRELIMINARY PAGES)
# ─────────────────────────────────────────────────────────────
list_fig_start = -1
for i, p in enumerate(doc.paragraphs[:75]):
    if p.text.strip() == "LIST OF FIGURES":
        list_fig_start = i
        break

if list_fig_start != -1:
    fig_titles = [
        "Figure 4.1: SOLAR SCAN Diagnostic Scan Lab — Light Studio Interface Mode",
        "Figure 4.2: SOLAR SCAN Diagnostic Scan Lab — True OLED Pitch Black Dark Mode",
        "Figure 4.3: Autonomous Drone Inspection GIS Module — Light Studio Mode",
        "Figure 4.4: Autonomous Drone Inspection GIS Module — True OLED Pitch Black Dark Mode",
        "Figure 4.5: Fleet Analytics & Defect Distribution Telemetry — Light Studio Mode",
        "Figure 4.6: Fleet Analytics & Defect Distribution Telemetry — True OLED Pitch Black Dark Mode",
        "Figure 4.7: Evidence Hub & Normalized 9x9 Confusion Matrix — Light Studio Mode",
        "Figure 4.8: Evidence Hub & Normalized 9x9 Confusion Matrix — True OLED Pitch Black Dark Mode"
    ]
    for k, title in enumerate(fig_titles):
        doc.paragraphs[list_fig_start + 1 + k].text = title
    print("Updated List of Figures in preliminary pages.")

# ─────────────────────────────────────────────────────────────
# 3. UPDATE SECTION 4.8 FUNCTIONAL MANUAL & UPGRADES
# ─────────────────────────────────────────────────────────────
sec_4_8_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "4.8 Application User Interface" in p.text:
        sec_4_8_idx = i
        break

print(f"Found Section 4.8 at P{sec_4_8_idx}")

# Update intro paragraph P+1
doc.paragraphs[sec_4_8_idx + 1].text = (
    "The SOLAR SCAN web application is organized into four primary operational workspaces accessible "
    "via the tactile hardware sidebar navigation on desktop and the bottom navigation bar on mobile. "
    "The interface features a dual-theme architecture supporting both Light Studio Mode (#FFFFFF / #F8FAFC) "
    "and True OLED Pitch Black Dark Mode (#000000 / #0A0A0A) for optimal visibility under bright sunlight or "
    "dark field conditions. The following describes the system's core functional components and recent upgrades:"
)

# Update descriptions in 4.8
doc.paragraphs[sec_4_8_idx + 2].text = (
    "Scan Lab Workspace & Dual-Engine Failover: The central diagnostic workspace where field technicians upload "
    "panel imagery via camera capture or local file upload. A tactile hardware engine toggle enables switching between "
    "the live Google Cloud Vision API for cloud-based annotations and the offline YOLOv8 TFLite Simulator for sub-180ms "
    "edge execution. An automatic Zero-Failure failover catches cloud API billing or network exceptions and seamlessly "
    "diverts inference to the local YOLOv8 engine without crashing or displaying disruptive error screens."
)

doc.paragraphs[sec_4_8_idx + 3].text = (
    "Dual-Theme Studio & Hardware Command System: An integrated theme engine featuring a crisp Light Studio Mode "
    "(#FFFFFF / #F8FAFC canvas with #0F172A hardware command controls) and a True OLED Pitch Black Dark Mode "
    "(#000000 pure black canvas with luminous #38BDF8 cyan and #FBBF24 amber telemetry). Technicians can toggle between "
    "modes with a single click in the sidebar or mobile header, with preferences automatically persisting in browser storage."
)

doc.paragraphs[sec_4_8_idx + 4].text = (
    "Live Internet-Connected AI Technical Assistant: A floating conversational AI companion capable of diagnosing "
    "photovoltaic issues in real time. The assistant connects directly to open technical encyclopedias and web knowledge "
    "APIs to look up real-time troubleshooting steps for inverter error codes, ground isolation faults (Riso low), and "
    "multimeter testing procedures, with optional multi-model Google Gemini reasoning (gemini-2.0-flash, gemini-1.5-flash)."
)

# ─────────────────────────────────────────────────────────────
# 4. CHIP SCREENSHOTS (BOTH LIGHT AND DARK MODE) INTO SECTION 4.8
# ─────────────────────────────────────────────────────────────
gallery_idx = -1
for i, p in enumerate(doc.paragraphs):
    if "Visual Interface Gallery" in p.text:
        gallery_idx = i
        break

print(f"Found Visual Interface Gallery at P{gallery_idx}")

# The 8 screenshots mapping
screenshot_mappings = [
    ("documents/screenshots/solarscan_light_scanlab.png", "Figure 4.1: SOLAR SCAN Diagnostic Scan Lab — Light Studio Interface Mode showing upload zone, engine switch, and defect preview."),
    ("documents/screenshots/solarscan_dark_scanlab.png", "Figure 4.2: SOLAR SCAN Diagnostic Scan Lab — True OLED Pitch Black Dark Mode featuring tactile hardware controls and high-contrast telemetry."),
    ("documents/screenshots/solarscan_light_dronemap.png", "Figure 4.3: Autonomous Drone Inspection GIS Module — Light Studio Mode displaying 24-array solar field flight grid and GPS panel status pins."),
    ("documents/screenshots/solarscan_dark_dronemap.png", "Figure 4.4: Autonomous Drone Inspection GIS Module — True OLED Pitch Black Dark Mode with live telemetry terminal and thermal IR spectrum."),
    ("documents/screenshots/solarscan_light_analytics.png", "Figure 4.5: Fleet Analytics & Defect Distribution Telemetry — Light Studio Mode showing health scores, loss ratios, and severity breakdowns."),
    ("documents/screenshots/solarscan_dark_analytics.png", "Figure 4.6: Fleet Analytics & Defect Distribution Telemetry — True OLED Pitch Black Dark Mode with luminous trend lines and anomaly charts."),
    ("documents/screenshots/solarscan_light_evidence.png", "Figure 4.7: Evidence Hub & Normalized 9x9 Confusion Matrix — Light Studio Mode presenting per-class precision and training convergence curves."),
    ("documents/screenshots/solarscan_dark_evidence.png", "Figure 4.8: Evidence Hub & Normalized 9x9 Confusion Matrix — True OLED Pitch Black Dark Mode with high-contrast matrix heatmaps.")
]

# Gallery starts at gallery_idx
# P+1 is intro text
# P+2 is Image 1, P+3 is Caption 1
# P+4 is Image 2, P+5 is Caption 2, etc.

for idx, (img_file, caption_text) in enumerate(screenshot_mappings):
    img_p_idx = gallery_idx + 2 + (idx * 2)
    cap_p_idx = img_p_idx + 1
    
    # 1. Update Image paragraph
    img_p = doc.paragraphs[img_p_idx]
    img_p._p.clear_content()
    img_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(img_file):
        run = img_p.add_run()
        run.add_picture(img_file, width=docx.shared.Inches(5.7))
        print(f"  [+] Embedded {img_file} at P{img_p_idx}")
    else:
        print(f"  [!] Missing image: {img_file}")
        
    # 2. Update Caption paragraph
    cap_p = doc.paragraphs[cap_p_idx]
    cap_p.text = caption_text
    cap_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    for r in cap_p.runs:
        r.font.size = docx.shared.Pt(10)
        r.font.italic = True

print("All 8 screenshots (both Dark Mode and Light Mode) chipped into Section 4.8!")

# ─────────────────────────────────────────────────────────────
# 5. SAVE TO ALL 4 REQUIRED LOCATIONS
# ─────────────────────────────────────────────────────────────
save_locations = [
    r"documents/Solar Scan Final Year Project Documentation.docx",
    r"C:\Users\amagy\Desktop\Solar Scan Final Year Project Documentation.docx",
    r"C:\Users\amagy\Downloads\Solar Scan Final Year Project Documentation.docx",
    r"C:\Users\amagy\Documents\Solar Scan Final Year Project Documentation.docx"
]

# Save primary
doc.save(save_locations[0])
print(f"Primary document saved: {save_locations[0]}")

# Copy to other 3 locations
for loc in save_locations[1:]:
    try:
        shutil.copy2(save_locations[0], loc)
        print(f"Successfully synced to: {loc}")
    except Exception as e:
        print(f"Warning: could not sync to {loc}: {e}")

print("Documentation update complete!")
