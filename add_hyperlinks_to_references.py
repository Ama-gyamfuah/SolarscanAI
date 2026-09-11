import docx
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from lxml import etree

# ─────────────────────────────────────────────────────────────
# HELPER: Add a hyperlink run inside a paragraph
# ─────────────────────────────────────────────────────────────
def add_hyperlink(paragraph, text, url):
    """
    Inserts a clickable hyperlink into a paragraph.
    Styled in blue, underlined — standard APA hyperlink format.
    """
    # Add relationship to document
    part = paragraph.part
    r_id = part.relate_to(url, 
        'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
        is_external=True)

    # Build <w:hyperlink> element
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), r_id)
    hyperlink.set(qn('w:history'), '1')

    # Build <w:r> run element inside hyperlink
    new_run = OxmlElement('w:r')

    # Run properties — blue, underlined, Times New Roman 12pt
    rPr = OxmlElement('w:rPr')

    rStyle = OxmlElement('w:rStyle')
    rStyle.set(qn('w:val'), 'Hyperlink')
    rPr.append(rStyle)

    color_el = OxmlElement('w:color')
    color_el.set(qn('w:val'), '0563C1')
    rPr.append(color_el)

    u_el = OxmlElement('w:u')
    u_el.set(qn('w:val'), 'single')
    rPr.append(u_el)

    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), 'Times New Roman')
    rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    rPr.append(rFonts)

    sz = OxmlElement('w:sz')
    sz.set(qn('w:val'), '24')  # 12pt = 24 half-points
    rPr.append(sz)

    new_run.append(rPr)

    # Text node
    t = OxmlElement('w:t')
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    return hyperlink


# ─────────────────────────────────────────────────────────────
# HELPER: Rebuild a reference paragraph with hyperlinks
# ─────────────────────────────────────────────────────────────
URL_REGEX = re.compile(r'(https?://[^\s]+)')

def rebuild_paragraph_with_links(para):
    """
    Given a paragraph containing plain-text URLs, rebuild it so
    each URL becomes a proper Word hyperlink. Returns True if changed.
    """
    full_text = para.text
    if not URL_REGEX.search(full_text):
        return False  # Nothing to do

    # Split the text into parts: [plain_text, url, plain_text, url, ...]
    parts = URL_REGEX.split(full_text)
    if len(parts) <= 1:
        return False

    # Remember formatting from first run
    font_name = 'Times New Roman'
    font_size_half_pt = '24'  # 12pt
    is_bold = False
    for run in para.runs:
        if run.text.strip():
            font_name = run.font.name or 'Times New Roman'
            font_size_half_pt = str(int((run.font.size.pt if run.font.size else 12) * 2))
            is_bold = run.bold or False
            break

    # Clear all existing runs from the paragraph XML
    p_xml = para._p
    for child in list(p_xml):
        tag = child.tag.split('}')[-1] if '}' in child.tag else child.tag
        if tag in ('r', 'hyperlink', 'bookmarkStart', 'bookmarkEnd'):
            p_xml.remove(child)

    # Rebuild: alternate plain text runs and hyperlink elements
    for part in parts:
        if not part:
            continue
        if URL_REGEX.match(part):
            # This is a URL — make it a hyperlink
            hyperlink_el = add_hyperlink(para, part, part)
            p_xml.append(hyperlink_el)
        else:
            # Plain text run
            run_el = OxmlElement('w:r')
            rPr = OxmlElement('w:rPr')

            rFonts = OxmlElement('w:rFonts')
            rFonts.set(qn('w:ascii'), font_name)
            rFonts.set(qn('w:hAnsi'), font_name)
            rPr.append(rFonts)

            sz = OxmlElement('w:sz')
            sz.set(qn('w:val'), font_size_half_pt)
            rPr.append(sz)

            if is_bold:
                rPr.append(OxmlElement('w:b'))

            run_el.append(rPr)

            t = OxmlElement('w:t')
            t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t.text = part
            run_el.append(t)

            p_xml.append(run_el)

    return True


# ─────────────────────────────────────────────────────────────
# MAIN: Load document, process references, save
# ─────────────────────────────────────────────────────────────
file_path = 'documents/Solar Scan Final Year Project Documentation.docx'
doc = Document(file_path)

print(f"Loaded document — {len(doc.paragraphs)} paragraphs.")

# Find REFERENCES heading
ref_start = None
for i, p in enumerate(doc.paragraphs):
    if p.text.strip().upper() == 'REFERENCES':
        ref_start = i
        print(f"REFERENCES heading found at paragraph {i}.")
        break

if ref_start is None:
    print("ERROR: REFERENCES heading not found!")
    exit(1)

# Process every paragraph from the REFERENCES heading onwards
total_links_added = 0
for i in range(ref_start + 1, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if not p.text.strip():
        continue
    changed = rebuild_paragraph_with_links(p)
    if changed:
        # Count URLs in the paragraph
        urls = URL_REGEX.findall(p.text)
        total_links_added += len(urls)
        print(f"  P{i}: {len(urls)} link(s) added — {p.text[:80]}...")

print(f"\nTotal hyperlinks added: {total_links_added}")

# ─────────────────────────────────────────────────────────────
# SAVE
# ─────────────────────────────────────────────────────────────
targets = [
    "documents/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Desktop/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Downloads/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Documents/Solar Scan Final Year Project Documentation.docx",
]
for t in targets:
    try:
        doc.save(t)
        print(f"Saved: {t}")
    except Exception as e:
        print(f"Error saving {t}: {e}")

import os
os.system('powershell -Command "Start-Process \'documents/Solar Scan Final Year Project Documentation.docx\'"')
print("\nDone! All reference URLs are now clickable hyperlinks in Word.")
