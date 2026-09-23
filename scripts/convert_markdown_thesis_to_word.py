import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

doc = docx.Document()

# Set Standard Margins
for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Read the extracted Markdown file
md_file_path = "C:/Users/amagy/.gemini/antigravity/brain/dd5ab81c-e9c6-4d54-856e-59de7e59fddb/solarscan_official_thesis_ch1_5.md"
with open(md_file_path, "r", encoding="utf-8") as f:
    lines = f.readlines()

in_code_block = False
code_buffer = []

for line in lines:
    text = line.strip()
    
    # Handle code blocks
    if text.startswith("```"):
        if in_code_block:
            # Output buffered code
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.right_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run("\n".join(code_buffer))
            run.font.name = 'Courier New'
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor(14, 116, 144) # Cyan/Dark Teal
            code_buffer = []
            in_code_block = False
        else:
            in_code_block = True
            code_buffer = []
        continue

    if in_code_block:
        code_buffer.append(line.rstrip())
        continue

    if not text:
        continue

    # Title & Main Headings
    if text.startswith("# CHAPTER") or text.startswith("CHAPTER"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text.replace("#", "").strip())
        r.bold = True
        r.font.size = Pt(14)
        r.font.name = 'Arial'
        r.font.color.rgb = RGBColor(28, 25, 23)
        pBrd = OxmlElement('w:pBrd')
        bottom = OxmlElement('w:bottom')
        bottom.set(qn('w:val'), 'single')
        bottom.set(qn('w:sz'), '18')
        bottom.set(qn('w:space'), '6')
        bottom.set(qn('w:color'), 'D97706') # Amber Gold
        pBrd.append(bottom)
        p._p.get_or_add_pPr().append(pBrd)

    elif text.startswith("## ") or text.startswith("# "):
        h_text = text.replace("##", "").replace("#", "").strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.name = 'Arial'
        r.font.color.rgb = RGBColor(217, 119, 6)

    elif text.startswith("### "):
        h_text = text.replace("###", "").strip()
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.name = 'Arial'
        r.font.color.rgb = RGBColor(68, 64, 60)

    elif text.startswith("- ") or text.startswith("* "):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(text[2:])
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(68, 64, 60)

    else:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(6)
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(68, 64, 60)

output_file = "documents/SolarScanAI_Full_9600Word_Thesis_Markdown_Converted.docx"
doc.save(output_file)
print(f"Converted full {len(lines)} lines markdown thesis into Word document successfully at {output_file}!")
