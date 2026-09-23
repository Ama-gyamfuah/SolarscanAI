import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os
import re

file_path = 'documents/Solar Scan Final Year Project Documentation.docx'
doc = docx.Document(file_path)

print(f"Loaded {file_path} with {len(doc.paragraphs)} paragraphs.")

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Add Chapter 4 Section 4.8: Complete Visual UI Tour & Button-by-Button Operational Manual
p_m_title = doc.add_paragraph()
p_m_title.paragraph_format.space_before = Pt(22)
p_m_title.paragraph_format.space_after = Pt(8)
p_m_title.paragraph_format.keep_with_next = True
r_mt = p_m_title.add_run("4.8 Comprehensive Web Application UI Tour & Button-by-Button Operation Manual")
r_mt.bold = True
r_mt.font.size = Pt(13)
r_mt.font.name = 'Arial'
r_mt.font.color.rgb = RGBColor(217, 119, 6)

p_intro = doc.add_paragraph()
p_intro.paragraph_format.line_spacing = 1.15
p_intro.paragraph_format.space_after = Pt(8)
p_intro.add_run(
    "This section provides an exhaustive visual tour and functional user manual for the entire SolarScan AI application interface. "
    "To ensure that any evaluator, supervisor, or field technician can navigate the software with total clarity, every screen tab, interactive button, "
    "input field, dropdown control, diagnostic trigger, system outcome, and operational consequence is documented in detail below."
)

# Subsection 4.8.1: System Navigation & Sidebar Workspace Switcher
p_s1 = doc.add_paragraph()
p_s1.paragraph_format.space_before = Pt(14)
p_s1.paragraph_format.space_after = Pt(4)
r_s1 = p_s1.add_run("4.8.1 Application Layout & Sidebar Navigation System")
r_s1.bold = True
r_s1.font.size = Pt(11)
r_s1.font.color.rgb = RGBColor(28, 25, 23)

p_s1_desc = doc.add_paragraph()
p_s1_desc.paragraph_format.line_spacing = 1.15
p_s1_desc.paragraph_format.space_after = Pt(6)
p_s1_desc.add_run(
    "The application layout features a persistent left-hand Sidebar Navigation (desktop) and Bottom Navigation Bar (mobile) "
    "that enables instant switching between five primary application modules:\n"
    "1. Scan Lab (`/scan` tab): The primary diagnostic workspace for uploading panel imagery, configuring engine overrides, running scans, and calculating financial ROI.\n"
    "2. Evidence Hub (`/evidence` tab): Displays academic benchmarks, 9x9 normalized confusion matrices, system architecture block diagrams, and literature citations.\n"
    "3. Analytics (`/analytics` tab): Provides historical scan distribution metrics, defect frequency charts, and severity breakdowns over time.\n"
    "4. Drone GIS Map (`/map` tab): Renders an interactive 2D solar farm array grid displaying individual panel status indicators (green = healthy, red = defect) with GPS coordinate popups.\n"
    "5. Developer Docs (`/docs` tab): Contains the university assessment checklist, Python/PyTorch model training code blocks, and the 10-question SUS usability evaluation panel."
)

# Table detailing all Sidebar Buttons & Outcomes
t_nav = doc.add_table(rows=6, cols=4)
t_nav.style = 'Table Grid'
th_nav = ["Navigation Button", "UI Route", "Trigger Function / Operation", "Outcome & System Consequence"]
for idx, h in enumerate(th_nav):
    c = t_nav.rows[0].cells[idx]
    c.paragraphs[0].add_run(h).bold = True
    c.paragraphs[0].runs[0].font.name = 'Arial'
    c.paragraphs[0].runs[0].font.size = Pt(9.5)
    c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_shading(c, "D97706")
    set_cell_margins(c)

nav_rows = [
    ("Scan Lab Button", "/scan", "Clicking tab icon", "Switches active viewport to main scan hub. Loads image uploader, engine toggles, and HUD canvas."),
    ("Evidence Hub Button", "/evidence", "Clicking tab icon", "Switches active viewport to evidence hub. Displays confusion matrices, per-class mAP tables, and citations."),
    ("Analytics Button", "/analytics", "Clicking tab icon", "Switches active viewport to analytics charts. Summarizes scan history logs and defect pie charts."),
    ("Drone GIS Map Button", "/map", "Clicking tab icon", "Switches active viewport to drone map. Loads interactive solar farm grid layout with GPS panel markers."),
    ("Developer Docs Button", "/docs", "Clicking tab icon", "Switches active viewport to documentation. Displays training scripts, copy-code triggers, and SUS evaluation scores.")
]

for row_idx, r_data in enumerate(nav_rows):
    row = t_nav.rows[row_idx + 1]
    for col_idx, text in enumerate(r_data):
        c = row.cells[col_idx]
        c.paragraphs[0].add_run(text)
        c.paragraphs[0].runs[0].font.name = 'Arial'
        c.paragraphs[0].runs[0].font.size = Pt(9)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(68, 64, 60)
        if row_idx % 2 == 0:
            set_cell_shading(c, "FAF9F6")
        set_cell_margins(c)

# Subsection 4.8.2: Scan Lab Controls & Buttons Manual
p_s2 = doc.add_paragraph()
p_s2.paragraph_format.space_before = Pt(14)
p_s2.paragraph_format.space_after = Pt(4)
r_s2 = p_s2.add_run("4.8.2 Scan Lab Interactive Buttons, Dropdowns & Controls Manual")
r_s2.bold = True
r_s2.font.size = Pt(11)
r_s2.font.color.rgb = RGBColor(28, 25, 23)

t_btns = doc.add_table(rows=9, cols=4)
t_btns.style = 'Table Grid'
th_btns = ["UI Button / Control", "Control Type", "User Action / Input", "Detailed Outcome & Consequence"]
for idx, h in enumerate(th_btns):
    c = t_btns.rows[0].cells[idx]
    c.paragraphs[0].add_run(h).bold = True
    c.paragraphs[0].runs[0].font.name = 'Arial'
    c.paragraphs[0].runs[0].font.size = Pt(9.5)
    c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_shading(c, "D97706")
    set_cell_margins(c)

btn_rows = [
    ("Dual-Engine Switcher: 'Live Cloud Vision API'", "Tab Button", "Click to select", "Switches classifier to Google Cloud Vision REST API. Connects to cloud servers for deep pixel-level annotation."),
    ("Dual-Engine Switcher: 'YOLOv8 Edge Simulator'", "Tab Button", "Click to select", "Switches classifier to local offline TFLite prototype interface modeling <200ms mobile CPU field execution."),
    ("Configure API Key Button", "Header Button", "Click to open drawer", "Opens modal drawer allowing user to input/update Google Cloud Vision API key string (`AIzaSy...`). Saved to browser `localStorage`."),
    ("Preset Sample Buttons (Hotspot, Crack, Soiling, Healthy)", "Action Buttons", "Click any sample", "Loads corresponding pre-configured panel scan image directly into memory for instant testing."),
    ("Configure Diagnostic Override", "Dropdown Select", "Select defect state", "Forces classifier to run under chosen state ('Broken Panel', 'Dusty Surface', 'Hotspot') regardless of filename for controlled testing."),
    ("Run Scan / Analyze Image", "Primary Action", "Click button", "Triggers image preprocessing, executes selected AI engine, calculates thermodynamic watt loss, and renders plain-English summary card."),
    ("Re-Scan Image", "Secondary Action", "Click button", "Re-executes diagnostic calculations on currently loaded image without requiring file re-upload."),
    ("Reset Scanner", "Secondary Action", "Click button", "Clears current image memory, resets override parameters, and returns workspace to initial upload state.")
]

for row_idx, r_data in enumerate(btn_rows):
    row = t_btns.rows[row_idx + 1]
    for col_idx, text in enumerate(r_data):
        c = row.cells[col_idx]
        c.paragraphs[0].add_run(text)
        c.paragraphs[0].runs[0].font.name = 'Arial'
        c.paragraphs[0].runs[0].font.size = Pt(9)
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(68, 64, 60)
        if row_idx % 2 == 0:
            set_cell_shading(c, "FAF9F6")
        set_cell_margins(c)

# Subsection 4.8.3: Diagnostic Results, Output Cards & Financial ROI Sliders
p_s3 = doc.add_paragraph()
p_s3.paragraph_format.space_before = Pt(14)
p_s3.paragraph_format.space_after = Pt(4)
r_s3 = p_s3.add_run("4.8.3 Diagnostic Output Panel & Financial ROI Sliders Manual")
r_s3.bold = True
r_s3.font.size = Pt(11)
r_s3.font.color.rgb = RGBColor(28, 25, 23)

p_s3_desc = doc.add_paragraph()
p_s3_desc.paragraph_format.line_spacing = 1.15
p_s3_desc.paragraph_format.space_after = Pt(6)
p_s3_desc.add_run(
    "When a scan is executed, the right-hand panel renders four detailed output modules:\n"
    "1. Plain-English Summary Card: Rendered at the very top of the output. Displays a color-coded title (e.g. 'Broken / Damaged Panel'), failure cause analysis, and step-by-step repair guidance (Norman, 2013).\n"
    "2. HUD Canvas Viewport: Renders bounding boxes around isolated defects, confidence percentages, and a 10x10 Grad-CAM attention grid showing AI focal points.\n"
    "3. Wafer Thermodynamic Telemetry: Calculates nominal output (400W), Loss Percentage (%), Actual Output (Watts), and Cell Temperature Rise (NOCT thermal equations; Duffie & Beckman, 2013).\n"
    "4. Interactive Financial ROI Calculator: Features range sliders for Solar Farm Panel Count (1 to 10,000 panels), Electricity Tariff ($/kWh), and Repair Cost ($). Adjusting these sliders dynamically recalculates Net Annual Revenue Loss ($) and displays an automated ROI justification score."
)

# Embed Application Screenshot Figures with Detailed Captions
app_figures = [
    ("Figure 4.5: SolarScan AI Web Dashboard Workspace — Live scan check for 'Broken / Damaged Panel' displaying bounding box localization, plain-English summary card, and yield loss telemetry.", "documents/fig_app_crack_scan.png"),
    ("Figure 4.6: Thermal IR Scan Viewport — Localized 'Overheating Spot' hotspot detection displaying 10x10 Grad-CAM attention grid overlay, cell temperature rise metrics, and junction box safety alerts.", "documents/fig_app_hotspot_scan.png"),
    ("Figure 4.7: Surface Soiling Workspace — 'Dirty / Dusty / Sandy' scan integrated with the Financial ROI Calculator, annual revenue loss slider ($), and cleaning cost analysis.", "documents/fig_app_soiling_scan.png"),
    ("Figure 4.8: Healthy Panel Baseline — Diagnostic check verifying 100% nominal operational capacity and zero power loss for perfect photovoltaic modules.", "documents/fig_app_healthy_scan.png")
]

for caption, img_path in app_figures:
    if os.path.exists(img_path):
        doc.add_paragraph().paragraph_format.space_before = Pt(12)
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(img_path, width=Inches(4.4))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(14)
        r_c = p_cap.add_run(caption)
        r_c.font.name = 'Arial'
        r_c.font.size = Pt(9)
        r_c.font.italic = True
        r_c.font.color.rgb = RGBColor(120, 113, 108)

# Save destination targets
targets = [
    "documents/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Desktop/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Downloads/Solar Scan Final Year Project Documentation.docx",
    "C:/Users/amagy/Documents/Solar Scan Final Year Project Documentation.docx",
    "Solar Scan Final Year Project Documentation.docx"
]

for t in targets:
    try:
        doc.save(t)
        print(f"Saved complete visual UI manual documentation to {t}")
    except Exception as e:
        print(f"Error saving to {t}: {e}")

# Launch Microsoft Word
try:
    os.system('powershell -Command "Start-Process \'documents/Solar Scan Final Year Project Documentation.docx\'"')
    print("Launched Microsoft Word with updated documentation!")
except Exception as e:
    print("Launch error:", e)
