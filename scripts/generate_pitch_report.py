import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

doc = docx.Document()

# Set standard margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

# Helper for cell shading
def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

# Helper for cell margins/padding
def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Title
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title_p.add_run("UNIVERSITY OF ENERGY AND NATURAL RESOURCES\n")
run.bold = True
run.font.size = Pt(11)
run.font.name = 'Arial'
run.font.color.rgb = RGBColor(120, 113, 108)

run_title = title_p.add_run("SOCIAL MEDIA PERFORMANCE & ELEVATOR PITCH ENGAGEMENT REPORT\n")
run_title.bold = True
run_title.font.size = Pt(14)
run_title.font.name = 'Arial'
run_title.font.color.rgb = RGBColor(28, 25, 23)

run_subtitle = title_p.add_run("Project Title: AI-Powered Solar Intelligence (SolarScan AI)\nSunyani, Bono Region, Ghana")
run_subtitle.font.size = Pt(10)
run_subtitle.font.italic = True
run_subtitle.font.name = 'Arial'
run_subtitle.font.color.rgb = RGBColor(87, 83, 78)

doc.add_paragraph().paragraph_format.space_after = Pt(12)

# Helper function to add a Section Heading
def add_section_header(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11)
    run.font.name = 'Arial'
    run.font.color.rgb = RGBColor(28, 25, 23)
    pBrd = OxmlElement('w:pBrd')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '12')
    bottom.set(qn('w:space'), '4')
    bottom.set(qn('w:color'), 'D97706') # Amber gold
    pBrd.append(bottom)
    p._p.get_or_add_pPr().append(pBrd)

# Executive Summary
add_section_header("1. EXECUTIVE SUMMARY")
p_exec = doc.add_paragraph()
p_exec.paragraph_format.line_spacing = 1.15
p_exec.paragraph_format.space_after = Pt(8)
run_exec = p_exec.add_run(
    "To support the presentation of the BSc IT Final Year Project, 'AI-Powered Solar Intelligence (SolarScan AI),' at the UENR Tech Fair 2026, an elevator pitch video was developed and posted across multiple social media platforms including TikTok, Instagram, and Facebook. The objective of this campaign was to showcase the practical applications of Computer Vision (YOLOv8) and Edge AI in solar panel maintenance, driving academic visibility and community engagement. Over the first 24 hours of release, the content garnered wide local interest, demonstrating high relative engagement, positive viewer sentiment, and peer support."
)
run_exec.font.name = 'Arial'
run_exec.font.size = Pt(10.5)
run_exec.font.color.rgb = RGBColor(68, 64, 60)

# Social Media Platforms Overview
add_section_header("2. CAMPAIGN METRICS DASHBOARD")
p_table_desc = doc.add_paragraph()
p_table_desc.paragraph_format.space_after = Pt(6)
r_table_desc = p_table_desc.add_run("The following table aggregates the quantitative reach and engagement metrics across the documented social media channels:")
r_table_desc.font.name = 'Arial'
r_table_desc.font.size = Pt(10)
r_table_desc.font.italic = True

table = doc.add_table(rows=5, cols=6)
table.style = 'Table Grid'

# Header row
headers = ["Platform", "Account Handle", "Content Type", "Reach / Views", "Likes", "Comments & Shares"]
for idx, h_text in enumerate(headers):
    cell = table.rows[0].cells[idx]
    cell.paragraphs[0].add_run(h_text).bold = True
    cell.paragraphs[0].runs[0].font.name = 'Arial'
    cell.paragraphs[0].runs[0].font.size = Pt(9.5)
    cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
    set_cell_shading(cell, "D97706") # Amber header
    set_cell_margins(cell, 100, 100, 120, 120)

rows_data = [
    ("TikTok / Short-Video", "gyamie_rose", "Video Reel / Pitch", "214 Views", "46 Likes", "11 Comments"),
    ("Instagram", "gyamie_", "Video Reel / Pitch", "N/A (Reel)", "12 Likes", "0 Comments, 1 Share"),
    ("Instagram (Poster)", "gyamie_", "Static Post (Poster)", "N/A (Feed)", "1 Like", "0 Comments"),
    ("Facebook", "Monica Gyamfuah", "Video / Post", "N/A (Post)", "4 Likes", "1 Comment, 1 Share")
]

for row_idx, r_data in enumerate(rows_data):
    row = table.rows[row_idx + 1]
    for col_idx, text in enumerate(r_data):
        cell = row.cells[col_idx]
        cell.paragraphs[0].add_run(text)
        cell.paragraphs[0].runs[0].font.name = 'Arial'
        cell.paragraphs[0].runs[0].font.size = Pt(9.5)
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(68, 64, 60)
        if row_idx % 2 == 0:
            set_cell_shading(cell, "FAF9F6") # Light tint
        set_cell_margins(cell, 80, 80, 120, 120)

# Set widths
col_widths = [Inches(1.5), Inches(1.2), Inches(1.3), Inches(1.0), Inches(0.8), Inches(1.2)]
for r in table.rows:
    for c_idx, width in enumerate(col_widths):
        r.cells[c_idx].width = width

# Total Stats callout
total_p = doc.add_paragraph()
total_p.paragraph_format.space_before = Pt(8)
total_run = total_p.add_run("TOTAL AGGREGATED METRICS: 214+ Views | 63 Likes | 12 Comments | 2 Shares")
total_run.bold = True
total_run.font.name = 'Arial'
total_run.font.size = Pt(10)
total_run.font.color.rgb = RGBColor(28, 25, 23)

# Qualitative Sentiment analysis
add_section_header("3. QUALITATIVE SENTIMENT & COMMENT ANALYSIS")
p_qual = doc.add_paragraph()
p_qual.paragraph_format.line_spacing = 1.15
p_qual.paragraph_format.space_after = Pt(6)
r_qual = p_qual.add_run("Analysis of the comment sections reveals highly positive reception, categorized into three distinct themes:")
r_qual.font.name = 'Arial'
r_qual.font.size = Pt(10)

themes = [
    ("1. Technical & Presentation Appreciation", 
     "Viewers specifically praised the quality of the presentation and the neatness of the execution. Comments like 'Good presentation 100% ❤️' (Bernard) and 'Neat presentation' (Rookie Senses) highlight that the audience found the delivery professional and clear."),
    ("2. Project Innovation & Expectation", 
     "Viewers expressed excitement about the system's potential and utility in real-world scenarios. The comment 'This will be nice' (Gloria) accompanied by a response from the creator 'Would be 😳' highlights the anticipation of seeing the completed tool in the industry."),
    ("3. Peer Support & Team Encouragement", 
     "As a Tech Fair presentation, there was massive peer support and encouragement from fellow students and friends. Comments like 'Go girlll 👍' (ladies_outlet), 'Amaaa ❤️❤️' (_Naryah), and friendly teasing local banter like 'Mop3 kasa dodo Monica' (Gilberto) showcase high community bonding and team morale.")
]

for title, desc in themes:
    p_t = doc.add_paragraph()
    p_t.paragraph_format.left_indent = Inches(0.2)
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run(title)
    r_t.bold = True
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(10)
    r_t.font.color.rgb = RGBColor(217, 119, 6)

    p_d = doc.add_paragraph()
    p_d.paragraph_format.left_indent = Inches(0.4)
    p_d.paragraph_format.space_after = Pt(6)
    r_d = p_d.add_run(desc)
    r_d.font.name = 'Arial'
    r_d.font.size = Pt(9.5)
    r_d.font.color.rgb = RGBColor(87, 83, 78)

# Visual Identity
add_section_header("4. PRESENTATION TEAM & BRAND IDENTITY")
p_brand = doc.add_paragraph()
p_brand.paragraph_format.line_spacing = 1.15
p_brand.paragraph_format.space_after = Pt(8)
r_brand = p_brand.add_run(
    "The video highlights a team of four (4) BSc IT Final Year students presenting at the campus of the University of Energy and Natural Resources (UENR). The presenters dressed professionally, reflecting the academic integrity of the event. The video features UENR Tech Fair 2026 branding overlays ('UENR Tech Fair 2026') in the top-left corner, ensuring brand compliance and establishing institutional association."
)
r_brand.font.name = 'Arial'
r_brand.font.size = Pt(10)
r_brand.font.color.rgb = RGBColor(68, 64, 60)

# Recommendations
add_section_header("5. ENGAGEMENT RECOMMENDATIONS FOR THE BOOTH")
p_recs_intro = doc.add_paragraph()
p_recs_intro.paragraph_format.space_after = Pt(4)
r_recs_intro = p_recs_intro.add_run("To capitalize on the social media momentum during the physical tech fair, the following strategies are recommended:")
r_recs_intro.font.name = 'Arial'
r_recs_intro.font.size = Pt(10)

recs = [
    ("Live Demo QR Codes", "Display a QR code leading to the hosted Vercel link or the local network URL on your booth poster so visitors can load the mobile app directly on their phones while watching the pitch."),
    ("Call to Action (CTA) updates", "Update the description of current social media posts to guide visitors directly to the physical booth location (e.g., 'Visit us at Booth #12 to try a live scan!')."),
    ("Leverage Local Banter for Engagement", "Respond actively to all comments with friendly invitations to try the scanner in person, turning virtual likes into physical booth traffic.")
]

for r_title, r_desc in recs:
    p_r = doc.add_paragraph(style='List Bullet')
    p_r.paragraph_format.space_after = Pt(4)
    run_rt = p_r.add_run(r_title + ": ")
    run_rt.bold = True
    run_rt.font.name = 'Arial'
    run_rt.font.size = Pt(9.5)
    run_rt.font.color.rgb = RGBColor(28, 25, 23)
    
    run_rd = p_r.add_run(r_desc)
    run_rd.font.name = 'Arial'
    run_rd.font.size = Pt(9.5)
    run_rd.font.color.rgb = RGBColor(68, 64, 60)

# Appendix Section for Images
add_section_header("6. APPENDIX: CAMPAIGN ENGAGEMENT EVIDENCE")
p_evidence = doc.add_paragraph()
p_evidence.paragraph_format.line_spacing = 1.15
p_evidence.paragraph_format.space_after = Pt(8)
r_evidence = p_evidence.add_run(
    "The screenshots below present the verified impressions, likes, shares, and comment threads "
    "documented from the live posts on Instagram, TikTok, and Facebook:"
)
r_evidence.font.name = 'Arial'
r_evidence.font.size = Pt(10)
r_evidence.font.color.rgb = RGBColor(68, 64, 60)

# Add pictures
try:
    import os
    proof_images = [
        ("Figure 1: Instagram Post and Reel Engagement", "documents/instagram_proof.jpg"),
        ("Figure 2: Facebook Pitch Video Post", "documents/facebook_proof.jpg"),
        ("Figure 3: TikTok Video Comments (Part 1)", "documents/tiktok_proof_1.jpg"),
        ("Figure 4: TikTok Video Comments (Part 2) & Feedbacks", "documents/tiktok_proof_2.jpg")
    ]
    for label, path in proof_images:
        if os.path.exists(path):
            doc.add_paragraph().paragraph_format.space_before = Pt(8)
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.add_run().add_picture(path, width=Inches(3.2))
            
            p_lbl = doc.add_paragraph()
            p_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_lbl.paragraph_format.space_after = Pt(12)
            r_lbl = p_lbl.add_run(label)
            r_lbl.font.name = 'Arial'
            r_lbl.font.size = Pt(9)
            r_lbl.font.italic = True
            r_lbl.font.color.rgb = RGBColor(120, 113, 108)
except Exception as e:
    print(f"Error adding images: {e}")

doc.save("documents/Social_Media_Performance_and_Pitch_Report.docx")
print("Social_Media_Performance_and_Pitch_Report.docx generated successfully!")
